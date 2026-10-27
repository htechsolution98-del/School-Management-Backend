import io
from datetime import datetime
from decimal import Decimal
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Flowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT

REPORT_COLUMNS_SPEC = {
    'sales': [
        ('Date', lambda r: str(r.get('issue_date') or r.get('date') or '')),
        ('Student Name', lambda r: str(r.get('student_name') or '')),
        ('Class & Div', lambda r: f"{r.get('class_name') or r.get('student_class') or ''} {r.get('student_division') or ''}".strip() or '-'),
        ('Item Name', lambda r: str(r.get('item_name') or '')),
        ('Item Code', lambda r: str(r.get('item_code') or '')),
        ('Variant', lambda r: f"{r.get('size') or r.get('size_name') or ''} / {r.get('color') or r.get('color_name') or ''}".strip(' /') or '-'),
        ('Qty', lambda r: str(r.get('quantity') or 0)),
        ('Charging Type', lambda r: str(r.get('charging_type') or '')),
        ('Unit Cost (₹)', lambda r: f"{float(r.get('unit_cost') or 0):.2f}"),
        ('Total Cost (₹)', lambda r: f"{float(r.get('total_cost') or 0):.2f}"),
        ('Revenue (₹)', lambda r: f"{float(r.get('charged_amount') or 0):.2f}"),
        ('Profit (₹)', lambda r: f"{float(r.get('gross_profit') or r.get('profit') or 0):.2f}"),
    ],
    'purchases': [
        ('Date', lambda r: str(r.get('purchase_date') or r.get('date') or '')),
        ('Invoice #', lambda r: str(r.get('invoice_number') or '')),
        ('Supplier', lambda r: str(r.get('supplier_name') or 'Direct / Internal')),
        ('Academic Year', lambda r: str(r.get('academic_year_name') or '')),
        ('Items Count', lambda r: str(len(r.get('items', [])) if isinstance(r.get('items'), list) else (r.get('items_count') or 1))),
        ('Subtotal (₹)', lambda r: f"{float(r.get('subtotal') or 0):.2f}"),
        ('Tax (₹)', lambda r: f"{float(r.get('tax_amount') or 0):.2f}"),
        ('Discount (₹)', lambda r: f"{float(r.get('discount_amount') or 0):.2f}"),
        ('Grand Total (₹)', lambda r: f"{float(r.get('grand_total') or 0):.2f}"),
        ('Status', lambda r: str(r.get('status') or 'CONFIRMED')),
    ],
    'stock': [
        ('Date', lambda r: str(r.get('date') or r.get('transaction_date') or '')),
        ('Item Name', lambda r: str(r.get('item_name') or '')),
        ('Category', lambda r: str(r.get('category') or '-')),
        ('Variant', lambda r: f"{r.get('size') or ''} / {r.get('color') or ''}".strip(' /') or '-'),
        ('Qty', lambda r: str(r.get('quantity') or 0)),
        ('Balance After', lambda r: str(r.get('balance_after') or 0)),
        ('Unit Cost (₹)', lambda r: f"{float(r.get('unit_cost') or 0):.2f}"),
        ('Total Value (₹)', lambda r: f"{float(r.get('recorded_value') or 0):.2f}"),
    ],
    'replacements': [
        ('Date', lambda r: str(r.get('date') or '')),
        ('Student Name', lambda r: str(r.get('student_name') or '')),
        ('Item Name', lambda r: str(r.get('item_name') or '')),
        ('Qty', lambda r: str(r.get('quantity') or 0)),
        ('Reason', lambda r: str(r.get('reason_display') or r.get('reason') or '')),
        ('Status', lambda r: str(r.get('status') or '')),
        ('Admin Remarks', lambda r: str(r.get('admin_remarks') or '-')),
    ],
    'issues_ledger': [
        ('Issue Date', lambda r: str(r.get('issue_date') or '')),
        ('Student Name', lambda r: str(r.get('student_name') or '')),
        ('Class & Div', lambda r: f"{r.get('student_class') or ''} {r.get('student_division') or ''}".strip() or '-'),
        ('Item Issued', lambda r: str(r.get('item_name') or '')),
        ('Variant', lambda r: f"{r.get('size_name') or ''} {r.get('color_name') or ''}".strip() or '-'),
        ('Quantity', lambda r: str(r.get('quantity') or 0)),
        ('Status', lambda r: str(r.get('status') or 'ISSUED')),
    ]
}


def get_column_specs(report_type):
    return REPORT_COLUMNS_SPEC.get(report_type, REPORT_COLUMNS_SPEC['sales'])


def export_inventory_report_excel(report_type, rows, school=None, title=None):
    """
    Generate an Excel (.xlsx) file with strictly aligned headers and row mapping.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (title or report_type).title()[:30]

    specs = get_column_specs(report_type)
    headers = [col[0] for col in specs]

    # Title Banner
    school_name = school.name if school and hasattr(school, 'name') else "School Inventory Management"
    report_title = f"{school_name} - {title or report_type.replace('_', ' ').title()} Report"

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1, value=report_title)
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="1E293B")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    subtitle_cell = ws.cell(row=2, column=1, value=f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Total Records: {len(rows)}")
    subtitle_cell.font = Font(name="Calibri", size=9, italic=True, color="64748B")
    subtitle_cell.alignment = Alignment(horizontal="center", vertical="center")

    # Header Row (Row 4)
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    ws.row_dimensions[4].height = 28

    # Data Rows (Row 5+)
    row_alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    data_font = Font(name="Calibri", size=10, color="1E293B")

    for r_idx, row_data in enumerate(rows, start=5):
        is_even = (r_idx % 2 == 0)
        for c_idx, (_, extractor) in enumerate(specs, start=1):
            val = extractor(row_data)
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if is_even:
                cell.fill = row_alt_fill

            # Align numbers right, text left
            if isinstance(val, (int, float)) or (isinstance(val, str) and val.replace('.', '', 1).isdigit()):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[r_idx].height = 20

    # Adjust column widths
    from openpyxl.utils import get_column_letter
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for r in range(4, len(rows) + 5):
            val = ws.cell(row=r, column=col_idx).value
            if val:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)


    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def export_inventory_report_pdf(report_type, rows, school=None, title=None):
    """
    Generate a PDF file using ReportLab with strictly aligned headers and row mapping.
    """
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        textColor=colors.HexColor('#1E293B'),
        alignment=TA_LEFT
    )
    meta_style = ParagraphStyle(
        'DocMeta',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=11,
        textColor=colors.HexColor('#64748B'),
        alignment=TA_LEFT
    )
    th_style = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.white,
        alignment=TA_CENTER
    )
    td_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor('#1E293B'),
        alignment=TA_LEFT
    )

    specs = get_column_specs(report_type)
    headers = [col[0] for col in specs]

    elements: list[Flowable] = []
    school_name = school.name if school and hasattr(school, 'name') else "School Inventory Management"
    elements.append(Paragraph(f"<b>{school_name}</b> - {title or report_type.replace('_', ' ').title()} Report", title_style))
    elements.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Total Records: {len(rows)}", meta_style))
    elements.append(Spacer(1, 4 * mm))

    # Table data
    table_data = []
    # Header row
    table_data.append([Paragraph(h, th_style) for h in headers])

    # Body rows
    for row_data in rows:
        row_cells = []
        for _, extractor in specs:
            val = extractor(row_data)
            row_cells.append(Paragraph(val, td_style))
        table_data.append(row_cells)


    # Calculate column widths to fit page width (A4 landscape width = 297mm - 20mm margin = 277mm)
    total_avail_width = 277 * mm
    col_width = total_avail_width / max(1, len(headers))
    col_widths = [col_width] * len(headers)

    t = Table(table_data, colWidths=col_widths, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4F46E5')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('LEFTPADDING', (0, 0), (-1, -1), 3),
        ('RIGHTPADDING', (0, 0), (-1, -1), 3),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
    ]))

    elements.append(t)
    doc.build(elements)
    buf.seek(0)
    return buf.getvalue()
